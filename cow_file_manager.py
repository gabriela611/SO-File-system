import os
import json
from datetime import datetime
import docx
import re
import zlib
import base64

class GestorArchivos:
    def __init__(self, ruta_archivo):
        self.ruta_archivo = ruta_archivo
        self.nombre_archivo = os.path.basename(ruta_archivo)
        self.directorio_base = "lavacamu_data"
        self.directorio_archivos = os.path.join(self.directorio_base, "archivos")
        self.directorio_versiones_base = os.path.join(self.directorio_base, "versiones")
        self.directorio_inodos = os.path.join(self.directorio_base, "inodos")
        self.directorio_logs = os.path.join(self.directorio_base, "logs")

        self.inodo_path = os.path.join(self.directorio_inodos, f"{self.nombre_archivo}.json")
        # Directorio para los bloques del archivo
        self.versiones_dir = os.path.join(self.directorio_versiones_base, self.nombre_archivo)
        self.log_path = os.path.join(self.directorio_logs, f"{self.nombre_archivo}.log")

        self.BLOQUE_TAM_MAX = 4096  # 4KB por bloque

        self._asegurar_estructura()
        self._cargar_o_crear_inodo()

        # Si es la primera vez que se abre, se crea un bloque inicial (primer_bloque)
        if "primer_bloque" not in self.inodo or not self.inodo["primer_bloque"]:
            primer_nombre = self._crear_nuevo_bloque(inicial=True)
            self.inodo["primer_bloque"] = primer_nombre
            self.inodo["fat"] = {primer_nombre: {"next": None, "usado": 0}}
            self._guardar_inodo()

        # Si es la primera vez que se abre el archivo, se crea la versión 0
        if not self.inodo.get("versiones"):
            self._crear_version0()

    def _asegurar_estructura(self):
        for d in [self.directorio_base, self.directorio_archivos,
                  self.directorio_versiones_base, self.directorio_inodos,
                  self.directorio_logs, self.versiones_dir]:
            os.makedirs(d, exist_ok=True)

    def _cargar_o_crear_inodo(self):
        if os.path.exists(self.inodo_path):
            with open(self.inodo_path, "r", encoding="utf-8") as f:
                self.inodo = json.load(f)
        else:
            # Estructura inicial del inodo con FAT y lista de versiones vacía
            self.inodo = {
                "nombre": self.nombre_archivo,
                "primer_bloque": "",
                "fat": {},
                "versiones": [],
                "current_version": -1
            }
            self._guardar_inodo()
        if "current_version" not in self.inodo:
            self.inodo["current_version"] = len(self.inodo["versiones"]) - 1
            self._guardar_inodo()
        self.current_version = self.inodo["current_version"]

    def _guardar_inodo(self):
        with open(self.inodo_path, "w", encoding="utf-8") as f:
            json.dump(self.inodo, f, indent=4)

    def _registrar_log(self, mensaje):
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with open(self.log_path, "a", encoding="utf-8") as f:
            f.write(f"[{timestamp}] {mensaje}\n")

    # Métodos de bloques y FAT (almacenados como archivos JSON en versiones_dir)
    def _ruta_bloque(self, nombre_bloque):
        return os.path.join(self.versiones_dir, nombre_bloque)

    def _crear_nuevo_bloque(self, inicial=False):
        """
        Crea un nuevo bloque con un nombre incremental.
        Si es inicial, se establece en block_0001.json.
        """
        bloques_existentes = [f for f in os.listdir(self.versiones_dir) if f.startswith("block_") and f.endswith(".json")]
        if inicial or not bloques_existentes:
            nuevo_nombre = "block_0001.json"
        else:
            bloques_existentes.sort()
            ultimo = bloques_existentes[-1]
            numero = int(ultimo.replace("block_", "").replace(".json", ""))
            nuevo_nombre = f"block_{numero+1:04d}.json"
        bloque = {"nombre": nuevo_nombre, "contenido": "", "usado": 0, "max": self.BLOQUE_TAM_MAX, "paginas": [], "archivo": self.nombre_archivo }
        self._guardar_bloque_actual(bloque)
        self.inodo["fat"][nuevo_nombre] = {"next": None, "usado": 0}
        self._guardar_inodo()
        return nuevo_nombre

    def _obtener_bloque(self, nombre_bloque):
        ruta = self._ruta_bloque(nombre_bloque)
        if os.path.exists(ruta):
            with open(ruta, "r", encoding="utf-8") as f:
                bloque = json.load(f)
            # Validación fuerte de identidad
            if bloque.get("archivo") != self.nombre_archivo:
                print(f"⚠️ El bloque {nombre_bloque} no pertenece a este documento.")
                return None
            return bloque
        return None

    def _guardar_bloque_actual(self, bloque):
        ruta = self._ruta_bloque(bloque["nombre"])
        with open(ruta, "w", encoding="utf-8") as f:
            json.dump(bloque, f, indent=4)

    def _obtener_bloque_actual(self):
        bloque_actual = self._obtener_bloque(self.inodo["primer_bloque"])
        # Recorremos la lista enlazada de bloques usando la FAT
        while self.inodo["fat"][bloque_actual["nombre"]]["next"] is not None:
            siguiente = self.inodo["fat"][bloque_actual["nombre"]]["next"]
            bloque_actual = self._obtener_bloque(siguiente)
        return bloque_actual

    # Métodos para versiones
    def _crear_version0(self):
        extension = os.path.splitext(self.ruta_archivo)[-1].lower()
        if extension == ".txt":
            with open(self.ruta_archivo, "r", encoding="utf-8") as f:
                contenido = f.read()
        elif extension == ".docx":
            try:
                doc = docx.Document(self.ruta_archivo)
                contenido = "\n".join([p.text for p in doc.paragraphs])
            except Exception as e:
                print(f"❌ Error leyendo DOCX: {e}")
                contenido = ""
        else:
            contenido = ""
            with open(self.ruta_archivo, "w", encoding="utf-8") as f:
                f.write("")
        self._guardar_version_en_bloque(contenido, "v0")
        self._actualizar_current_version(len(self.inodo["versiones"]) - 1)
        print("✅ Versión 0 creada.")

    def _guardar_version_en_bloque(self, contenido, version_id):
        contenido_marcado = f"<v-{version_id}>{contenido}</v-{version_id}>"

        contenido_bytes = contenido_marcado.encode('utf-8')
        contenido_comprimido = zlib.compress(contenido_bytes)
        contenido_codificado = base64.b64encode(contenido_comprimido).decode('utf-8')
        longitud_total = len(contenido_codificado)

        offset_inicio = 0
        bloques_utilizados = []

        while offset_inicio < longitud_total:
            bloques_existentes = sorted([f for f in os.listdir(self.versiones_dir) if f.startswith("block_")])
            bloque_usado = None

            for nombre in bloques_existentes:
                bloque = self._obtener_bloque(nombre)
                if bloque and (bloque["max"] - bloque["usado"] >= 1):  # puede recibir algo
                    bloque_usado = bloque
                    break

            if not bloque_usado:
                nuevo_nombre = self._crear_nuevo_bloque()
                bloque_usado = self._obtener_bloque(nuevo_nombre)

            espacio_disponible = bloque_usado["max"] - bloque_usado["usado"]
            bytes_a_escribir = min(espacio_disponible, longitud_total - offset_inicio)
            fragmento = contenido_codificado[offset_inicio: offset_inicio + bytes_a_escribir]

            offset_bloque = bloque_usado["usado"]
            bloque_usado["contenido"] += fragmento
            bloque_usado["usado"] += bytes_a_escribir

            bloque_usado["paginas"].append({
                "version_id": version_id,
                "offset": offset_bloque,
                "longitud": bytes_a_escribir
            })

            # Ligadura fuerte al documento actual (ver mejora 2 abajo)
            bloque_usado["archivo"] = self.nombre_archivo

            self._guardar_bloque_actual(bloque_usado)

            bloques_utilizados.append({
                "bloque": bloque_usado["nombre"],
                "offset_inicio": offset_bloque,
                "offset_fin": offset_bloque + bytes_a_escribir
            })

            offset_inicio += bytes_a_escribir

        # Guardar metadata en el inodo
        version_metadata = {
            "id": version_id,
            "bloques": bloques_utilizados,
            "timestamp": datetime.now().strftime("%Y%m%d_%H%M%S")
        }

        self.inodo["versiones"].append(version_metadata)
        self._guardar_inodo()

    def _reconstruir_contenido(self, version_index):
        if version_index < 0 or version_index >= len(self.inodo["versiones"]):
            print("⚠️ Versión fuera de rango.")
            return ""

        version = self.inodo["versiones"][version_index]
        contenido_total = ""

        for bloque_info in version["bloques"]:
            bloque = self._obtener_bloque(bloque_info["bloque"])
            if not bloque:
                print(f"⚠️ No se pudo leer el bloque {bloque_info['bloque']}")
                continue

            inicio = bloque_info["offset_inicio"]
            fin = bloque_info["offset_fin"]
            fragmento = bloque["contenido"][inicio:fin]
            contenido_total += fragmento

        try:
            contenido_comprimido = base64.b64decode(contenido_total.encode('utf-8'))
            contenido_marcado = zlib.decompress(contenido_comprimido).decode('utf-8')
        except Exception as e:
            print("❌ Error al descomprimir:", e)
            return ""

        # Extraer el contenido real eliminando los marcadores
        patron = re.compile(r"<v-.*?>(.*?)</v-.*?>", re.DOTALL)
        resultado = patron.findall(contenido_marcado)
        return resultado[0] if resultado else ""

    def _actualizar_current_version(self, new_index):
        self.inodo["current_version"] = new_index
        self.current_version = new_index
        self._guardar_inodo()

    # Métodos públicos de lectura y escritura
    def leer(self):
        if self.current_version < 0 or self.current_version >= len(self.inodo["versiones"]):
            print("Versión actual fuera de rango")
            return

        version_info = self.inodo["versiones"][self.current_version]
        print(f"🔖 Metadatos de la versión {self.current_version}:\n")
        print(json.dumps(version_info, indent=4))

        contenido = self._reconstruir_contenido(self.current_version)
        if contenido:
            print(f"\n📄 Contenido (versión {self.current_version}):\n{contenido}")
        else:
            print("⚠️ No se pudo reconstruir el contenido.")

    def escribir(self, contenido):
        extension = os.path.splitext(self.ruta_archivo)[-1].lower()
        try:
            if extension == ".txt":
                with open(self.ruta_archivo, "a", encoding="utf-8") as f:
                    f.write(contenido + "\n")
            elif extension == ".docx":
                doc = docx.Document(self.ruta_archivo)
                doc.add_paragraph(contenido)
                doc.save(self.ruta_archivo)
            else:
                print("❌ Tipo de archivo no soportado para escritura.")
                return

            # Guardar la nueva versión
            nueva_version = f"v{len(self.inodo['versiones'])}"
            self._guardar_version_en_bloque(contenido, nueva_version)
            self._actualizar_current_version(len(self.inodo["versiones"]) - 1)
            self._registrar_log(f"Se escribió y guardó una nueva versión: {nueva_version}")
            print(f"✅ Contenido guardado en versión {nueva_version}.")

        except Exception as e:
            print(f"❌ Error al escribir contenido: {e}")

    def listar_versiones(self):
        versiones = self.inodo.get("versiones", [])

        if versiones:
            print("📜 Versiones registradas:")
            for i, version in enumerate(versiones):
                vid = version.get("id", "Sin ID")
                bloques = len(version.get("bloques", []))
                timestamp = version.get("timestamp", "Sin timestamp")
                tam = sum(b["offset_fin"] - b["offset_inicio"] for b in version["bloques"])
                print(f"{i + 1}. ID: {vid} | Bloques: {bloques} | Tamaño: {tam} bytes | Timestamp: {timestamp}")


            print(f"\n🔄 Versión actual: {self.current_version + 1}")
        else:
            print("⚠️ No hay versiones guardadas.")

    def rollback_backward(self):
        if self.current_version > 0:
            self._actualizar_current_version(self.current_version - 1)
            self._guardar_inodo()
            self._registrar_log(f"⏪ Retrocedió a versión v{self.current_version}")
            print(f"⏪ Retrocediste a la versión v{self.current_version}")
        else:
            print("🚫 Ya estás en la versión inicial. No se puede retroceder más.")

    def rollback_forward(self):
        if self.current_version < len(self.inodo["versiones"]) - 1:
            self._actualizar_current_version(self.current_version + 1)
            self._guardar_inodo()
            self._registrar_log(f"⏩ Avanzó a versión v{self.current_version}")

            print(f"⏩ Avanzaste a la versión v{self.current_version}")
        else:
            print("🚫 Ya estás en la versión más reciente. No se puede avanzar más.")

    # Funciones públicas para recuperación de versiones anteriores

    def ver_versiones_guardadas(self):
        print("📜 Versiones guardadas:")
        for i, version in enumerate(self.inodo["versiones"]):
            contenido = self._reconstruir_contenido(i)
            print(f"\n🔢 Versión ID: {version['id']}\n📄 Contenido:\n{contenido}")

    def restaurar_version(self, id_version_a_restaurar):
        # Buscar el índice de la versión indicada
        indices = [i for i, v in enumerate(self.inodo["versiones"]) if v["id"] == id_version_a_restaurar]
        if not indices:
            print(f"❌ La versión con ID {id_version_a_restaurar} no existe.")
            return

        index = indices[0]
        if self.inodo["current_version"] == index:
            print("⚠️ Ya estás en esa versión, no es necesario restaurarla.")
            return
        contenido_recuperado = self._reconstruir_contenido(index)
        nuevo_id = self._generar_nuevo_id()
        self._guardar_version_en_bloque(contenido_recuperado, nuevo_id)
        self._actualizar_current_version(len(self.inodo["versiones"]) - 1)
        self._registrar_log(f"Se restauró la versión {id_version_a_restaurar} como nueva versión {nuevo_id}")
        print(f"✅ Versión {id_version_a_restaurar} restaurada como nueva versión con ID: {nuevo_id}")

    def _generar_nuevo_id(self):
        # Genera un ID nuevo basado en los ya existentes
        ids = []
        for v in self.inodo["versiones"]:
            try:
                # Se asume que los IDs tienen el formato "vN" o numérico
                if isinstance(v["id"], str) and v["id"].startswith("v"):
                    ids.append(int(v["id"][1:]))
                else:
                    ids.append(int(v["id"]))
            except:
                continue
        ultimo = max(ids) if ids else 0
        return f"v{ultimo + 1}"

    def recuperar_version(self, version_id):
        # Buscar la versión por ID
        for index, version in enumerate(self.inodo["versiones"]):
            if version["id"] == version_id:
                self._actualizar_current_version(index)
                print(f"✅ Versión {version_id} restaurada correctamente.")
                return
        print(f"❌ Versión {version_id} no encontrada.")

    def ver_bloques_usados(self):
        print("🔍 Bloques por versión:")
        for v in self.inodo["versiones"]:
            print(f"{v['id']}: {[b['bloque'] for b in v['bloques']]}")

